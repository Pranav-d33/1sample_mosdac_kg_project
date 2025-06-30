import os
import json
import pdfplumber
import docx
import openpyxl

DOCS_DIR = "data/docs"
OUTPUT_FILE = "outputs/cleaned_json/cleaned_docs.json"

def extract_text_from_pdf(path):
    try:
        with pdfplumber.open(path) as pdf:
            return "\n".join(
                page.extract_text() or "" for page in pdf.pages
            )
    except Exception as e:
        print(f"❌ PDF extract failed: {path}: {e}")
        return ""

def extract_text_from_docx(path):
    try:
        doc = docx.Document(path)
        return "\n".join([p.text for p in doc.paragraphs])
    except Exception as e:
        print(f"❌ DOCX extract failed: {path}: {e}")
        return ""

def extract_text_from_xlsx(path):
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
        content = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                row_text = " | ".join([str(cell) if cell else "" for cell in row])
                content.append(row_text)
        return "\n".join(content)
    except Exception as e:
        print(f"❌ XLSX extract failed: {path}: {e}")
        return ""

def extract_all():
    os.makedirs(os.path.dirname(OUTPUT_FILE), exist_ok=True)
    data = []

    for fname in os.listdir(DOCS_DIR):
        fpath = os.path.join(DOCS_DIR, fname)
        if fname.lower().endswith(".pdf"):
            text = extract_text_from_pdf(fpath)
            source = "pdf"
        elif fname.lower().endswith(".docx"):
            text = extract_text_from_docx(fpath)
            source = "docx"
        elif fname.lower().endswith(".xlsx"):
            text = extract_text_from_xlsx(fpath)
            source = "xlsx"
        else:
            continue

        if text.strip():
            data.append({
                "filename": fname,
                "source": source,
                "text": text.strip()
            })

    with open(OUTPUT_FILE, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2, ensure_ascii=False)

    print(f"✅ Extracted text from {len(data)} documents → {OUTPUT_FILE}")

if __name__ == "__main__":
    extract_all()
